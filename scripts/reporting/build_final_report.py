"""Build the final Markdown and Word report from frozen Stage 1-7 outputs."""

from __future__ import annotations

import csv
import json
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[2]
MARKDOWN_PATH = PROJECT_ROOT / "docs" / "final_report.md"
DOCX_PATH = PROJECT_ROOT / "outputs" / "reports" / "STA4003_Final_Report.docx"
PDF_PATH = PROJECT_ROOT / "outputs" / "reports" / "STA4003_Final_Report.pdf"
STAGE7_SUMMARY_PATH = PROJECT_ROOT / "outputs" / "diagnostics" / "stage7_summary.json"
WORD_PATHS = (
    Path(r"C:\Program Files\Microsoft Office\root\Office16\WINWORD.EXE"),
    Path(r"C:\Program Files (x86)\Microsoft Office\root\Office16\WINWORD.EXE"),
)


def read_rows(relative_path: str) -> list[dict[str, str]]:
    with (PROJECT_ROOT / relative_path).open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def one_row(rows: list[dict[str, str]], **filters: str) -> dict[str, str]:
    matches = [row for row in rows if all(row.get(key) == value for key, value in filters.items())]
    if len(matches) != 1:
        raise ValueError(f"Expected one record for {filters}, found {len(matches)}")
    return matches[0]


def fnum(value: str) -> float:
    return float(value)


def fmt_level(value: str | float) -> str:
    return f"{float(value):,.2f}"


def fmt_mase(value: str | float) -> str:
    return f"{float(value):.6f}"


def fmt_change(value: str | float, digits: int = 2) -> str:
    return f"{float(value):+,.{digits}f}"


def stage_metric(path: str, model: str, *, horizon: int = 2, scope: str = "all") -> dict[str, str]:
    return one_row(read_rows(path), model=model, horizon=str(horizon), scope=scope)


def markdown_table(headers: list[str], rows: list[list[Any]]) -> str:
    def cell(value: Any) -> str:
        return str(value).replace("|", "/").replace("\n", " ")

    return "\n".join(
        [
            "| " + " | ".join(cell(value) for value in headers) + " |",
            "| " + " | ".join("---" for _ in headers) + " |",
            *["| " + " | ".join(cell(value) for value in row) + " |" for row in rows],
        ]
    )


def build_markdown() -> str:
    canonical = read_rows("data/processed/analysis_quarterly.csv")
    if not canonical:
        raise ValueError("Canonical analysis data is empty")
    data_start = canonical[0]["quarter"]
    data_end = canonical[-1]["quarter"]
    observed_pmi = [row["quarter"] for row in canonical if row.get("pmi", "")]
    pmi_start = observed_pmi[0]
    pmi_end = observed_pmi[-1]
    n_pmi = len(observed_pmi)

    results = read_rows("outputs/tables/final_results_summary.csv")
    primary_rows = [row for row in results if row["study/sample"] == "Primary 2005Q1-2026Q2"]
    pmi_result = one_row(results, **{"study/sample": "Stage 6A secondary 2010Q4-2026Q2", "model": "SARIMA+PMI"})
    sarima_primary = one_row(primary_rows, model="SARIMA")
    cny_primary = one_row(primary_rows, model="SARIMA+CNY")
    ets_primary = one_row(primary_rows, model="ETS")
    simple_models = [row for row in primary_rows if row["model"] in ("Historical Mean", "Naive", "Seasonal Naive")]

    selected_rows = read_rows("outputs/tables/sarima_origin_selection.csv")
    d1_count = sum(int(row["selected_D"]) == 1 for row in selected_rows)
    stage4 = __import__("json").loads((PROJECT_ROOT / "outputs/diagnostics/stage4_summary.json").read_text(encoding="utf-8"))
    full_order = stage4["full_sample_descriptive_selection"]
    seasonal_components = full_order["selected_seasonal_order"].strip("()").split(",")
    order_label = (
        f"SARIMA{full_order['selected_order']}"
        f"({','.join(seasonal_components[:3])})[{seasonal_components[3]}]"
    )

    seasonal = read_rows("outputs/tables/seasonal_summary.csv")
    q1_seasonal = one_row(seasonal, quarter_label="Q1")
    q2_seasonal = one_row(seasonal, quarter_label="Q2")
    q3_seasonal = one_row(seasonal, quarter_label="Q3")
    q4_seasonal = one_row(seasonal, quarter_label="Q4")
    complete_years = int(q4_seasonal["n_complete_years"])
    q4_above = int(round(fnum(q4_seasonal["share_complete_years_positive"]) * complete_years))
    q1_below = complete_years - int(round(fnum(q1_seasonal["share_complete_years_positive"]) * complete_years))
    q2_below = complete_years - int(round(fnum(q2_seasonal["share_complete_years_positive"]) * complete_years))

    stationarity = read_rows("outputs/tables/stationarity_tests.csv")
    level_test = one_row(stationarity, transformation="log_retail")
    diff_test = one_row(stationarity, transformation="first_difference")

    cny_incremental = one_row(
        read_rows("outputs/tables/cny_incremental_comparison.csv"), horizon="2", scope="all"
    )
    cny_q1 = one_row(read_rows("outputs/tables/cny_metrics.csv"), model="SARIMA+CNY", horizon="2", scope="q1")
    sarima_q1 = one_row(read_rows("outputs/tables/sarima_metrics.csv"), model="SARIMA", horizon="2", scope="q1")
    cny_shocks = read_rows("outputs/tables/cny_shock_sensitivity.csv")
    shock_full = one_row(cny_shocks, horizon="2", scope="all", sample="full")
    shock_excluded = one_row(cny_shocks, horizon="2", scope="all", sample="exclude_shocks")
    shock_targets = __import__("json").loads(
        (PROJECT_ROOT / "outputs/diagnostics/stage6b_summary.json").read_text(encoding="utf-8")
    )["shock_sensitivity"]["preset_targets"]

    pmi_delta = one_row(
        read_rows("outputs/tables/pmi_incremental_comparison.csv"),
        horizon="2",
        scope="all",
        reference_model="SARIMA-common",
        extended_model="SARIMA+PMI",
    )

    final_forecasts = read_rows("outputs/forecasts/final_forecast.csv")
    if [row["target"] for row in final_forecasts] != ["2026Q3", "2026Q4"]:
        raise ValueError("Final forecast target labels must be 2026Q3 and 2026Q4")
    forecast_table = []
    for row in final_forecasts:
        forecast_table.append(
            [
                row["target"],
                f"h={row['horizon']}",
                fmt_level(row["forecast_level_median"]),
                fmt_level(row["forecast_level_mean"]),
                f"[{fmt_level(row['lower_80'])}, {fmt_level(row['upper_80'])}]",
                f"[{fmt_level(row['lower_95'])}, {fmt_level(row['upper_95'])}]",
            ]
        )

    supplementary = read_rows("outputs/tables/final_forecast_supplementary.csv")
    supp_models = ("SARIMA", "SARIMA+CNY", "ETS")
    supplementary_table = []
    for model in supp_models:
        q3 = one_row(supplementary, model=model, target="2026Q3")
        q4 = one_row(supplementary, model=model, target="2026Q4")
        convention = q3["forecast_convention"]
        if model == "SARIMA+CNY":
            convention += "; future regressor is zero"
        supplementary_table.append([model, fmt_level(q3["forecast_level"]), fmt_level(q4["forecast_level"]), convention])

    diag = one_row(read_rows("outputs/tables/final_model_diagnostics.csv"), model="SARIMA")
    diagnostic_note = (
        f"The final full-sample fit converged ({diag['converged']}) with AIC {fnum(diag['AIC']):.3f}, "
        f"AICc {fnum(diag['AICc']):.3f}, and lag-8 Ljung-Box p-value {fnum(diag['ljung_box_pvalue']):.3f}. "
        "The complete parameter and residual diagnostics are retained in "
        "`outputs/tables/final_model_diagnostics.csv`."
    )

    historical_table = markdown_table(
        ["Model", "N", "MAE", "RMSE", "MASE"],
        [
            [row["model"], row["n"], fmt_level(row["MAE"]), fmt_level(row["RMSE"]), fmt_mase(row["MASE"])]
            for row in primary_rows
        ],
    )
    pmi_table = markdown_table(
        ["Model", "N", "MAE", "RMSE", "MASE"],
        [
            [model, record["n"], fmt_level(record["MAE"]), fmt_level(record["RMSE"]), fmt_mase(record["MASE"])]
            for model, record in (
                ("SARIMA-common", stage_metric("outputs/tables/pmi_metrics.csv", "SARIMA-common")),
                ("SARIMA+PMI", pmi_result),
            )
        ],
    )
    cny_table = markdown_table(
        ["Comparison", "N", "Change in MAE", "Change in RMSE", "Change in MASE"],
        [
            ["CNY minus SARIMA; h=2 all", cny_incremental["n"], fmt_change(cny_incremental["delta_MAE"]), fmt_change(cny_incremental["delta_RMSE"]), fmt_change(cny_incremental["delta_MASE"], 6)],
            ["CNY minus SARIMA; h=2 Q1", cny_q1["n"], fmt_change(fnum(cny_q1["MAE"]) - fnum(sarima_q1["MAE"])), fmt_change(fnum(cny_q1["RMSE"]) - fnum(sarima_q1["RMSE"])), fmt_change(fnum(cny_q1["MASE"]) - fnum(sarima_q1["MASE"]), 6)],
        ],
    )
    shock_table = markdown_table(
        ["Score sample", "N", "Change in MAE", "Change in RMSE", "Change in MASE"],
        [
            ["All preset targets", shock_full["n_sarima"], fmt_change(shock_full["delta_MAE"]), fmt_change(shock_full["delta_RMSE"]), fmt_change(shock_full["delta_MASE"], 6)],
            ["Exclude " + " and ".join(shock_targets), shock_excluded["n_sarima"], fmt_change(shock_excluded["delta_MAE"]), fmt_change(shock_excluded["delta_RMSE"]), fmt_change(shock_excluded["delta_MASE"], 6)],
        ],
    )
    forecast_markdown_table = markdown_table(
        ["Target", "Horizon", "Median", "Conditional mean point forecast", "80% prediction interval", "95% prediction interval"],
        forecast_table,
    )
    supplementary_markdown_table = markdown_table(
        ["Model", "2026Q3 forecast", "2026Q4 forecast", "Forecast convention"],
        supplementary_table,
    )

    lower_cny_mae = fnum(cny_incremental["delta_MAE"])
    lower_cny_rmse = fnum(cny_incremental["delta_RMSE"])
    lower_cny_mase = fnum(cny_incremental["delta_MASE"])
    lower_pmi_mae = fnum(pmi_delta["percent_change_MAE"])
    lower_pmi_rmse = fnum(pmi_delta["percent_change_RMSE"])
    lower_pmi_mase = fnum(pmi_delta["percent_change_MASE"])

    body = f"""# Forecasting Quarterly Retail Sales in China

**Course**: STA4003 Time Series Project

**Report**: Final forecast and research report

**Data through**: {data_end}

**Forecast origin**: 2026Q2

**Report date**: 2026-10-03

<!-- PAGE_BREAK -->

## 1 Abstract

This project evaluates whether Lunar New Year timing adds out-of-sample forecasting value beyond quarterly time-series structure in China's retail sales. Under the prespecified rolling-origin protocol, the frozen seasonal ARIMA baseline is strong: its h=2 all-target errors are lower than the three simple benchmarks and the fixed ETS comparator. SARIMA+CNY has only very small and mixed improvements, and its MAE and MASE advantages disappear when the {len(shock_targets)} preset shock targets are excluded. The fixed lag-2 PMI extension performs worse than its common SARIMA reference on all three h=2 metrics in its separate secondary window. A parsimonious plain SARIMA is therefore retained for the final forecast. It is refit on all {len(canonical)} observed quarters from {data_start} through {data_end}, then forecasts {final_forecasts[0]['target']} and {final_forecasts[1]['target']} from origin 2026Q2.

## 2 Research Question

The research question is whether deterministic information about Lunar New Year timing improves quarterly retail-sales forecasts beyond standard seasonal time-series dynamics. Forecast comparisons use the information available at each chronological origin. Coefficients and forecast differences are predictive descriptions under the specified models; they do not identify causal effects.

## 3 Data

The canonical sample contains {len(canonical)} quarterly observations from {data_start} through {data_end}. Retail sales are national aggregate amounts from the archived National Bureau of Statistics series and remain in RMB 100 million (亿元). The response used for model fitting is the natural logarithm of this level. The manufacturing PMI has {n_pmi} quarterly observations from {pmi_start} through {pmi_end}; its later forecasting test is kept on a shorter, explicitly separate window.

## 4 Data Construction

The quarterly retail series is constructed from monthly source data. For 1994-1999, quarterly values are sums of the available current-period monthly observations. From 2000 onward, quarterly values are differences of the cumulative series, which telescope to the December cumulative amount by construction. No economic values are imputed. The CNY date and position fraction are deterministic calendar metadata; the Stage 5 regressor uses only a training-centered Q1 fraction. PMI is aggregated as the arithmetic mean of its three monthly observations, then lagged before the Stage 6A evaluation.

An early descriptive diagnostic compares February retail with January retail. It is a February-to-January ratio, not February's share of combined January-February retail. The source-vintage differences and expected reporting gaps are documented in the Stage 1 data audit.

## 5 Exploratory Analysis and Stationarity

Retail sales rise strongly over the sample, while the logarithm makes proportional changes easier to compare and reduces the scale growth. The log transform does not itself establish stationarity. Across {complete_years} complete years, Q4 was above its annual mean in {q4_above} years; Q1 was below its annual mean in {q1_below} years and Q2 in {q2_below} years. These are descriptive within-year patterns, not a separate seasonal forecasting model.

![Seasonal subseries of log retail](../outputs/figures/seasonal_subseries_log.png)

*Figure 1. Annual-mean-centered log retail by quarter, showing the quarterly seasonal pattern. Source: Stage 2 frozen output.*

The prespecified ADF and KPSS tests on log levels point to non-stationarity under their stated deterministic terms (ADF p={fnum(level_test['adf_pvalue']):.3f}; KPSS probability is reported at its lower table bound). The first-difference tests are inconclusive at 5% (ADF p={fnum(diff_test['adf_pvalue']):.3f}; KPSS p={fnum(diff_test['kpss_pvalue']):.3f}). Stage 2 treated regular differencing as a cautious starting point and did not establish seasonal differencing. In Stage 4, rolling AICc selection chose D=1 at {d1_count} of {len(selected_rows)} origins, consistent with retaining D=0 in the dominant orders. Full test details are in `outputs/tables/stationarity_tests.csv` and `outputs/tables/sarima_origin_selection.csv`.

## 6 Forecast Evaluation Design

The primary rolling evaluation uses the common 2005Q1-2026Q2 target window and expanding training samples beginning in 1994Q1. For each target, the forecast origin is the target quarter minus the horizon. The primary endpoint is h=2 across all targets; h=1 is secondary, and Q1 is a prespecified subgroup. All models use the same original retail-level errors and origin-specific, training-only seasonal period-4 MASE scales. Lower MAE, RMSE, and MASE indicate lower errors.

The Stage 6A PMI extension uses its own secondary 2010Q4-2026Q2 target window because complete lag-2 PMI coverage starts later. Its SARIMA-common reference is refit on the same restricted training sample as SARIMA+PMI. The resulting sample is not combined with or substituted for the {sarima_primary['n']}-target primary window.

## 7 SARIMA Baseline

Stage 4 selected orders separately at each rolling origin using only that origin's training sample and AICc. The fixed candidate family had {int(full_order['candidate_count'])} candidates, and the chosen full-sample descriptive order was {order_label}. The rolling evaluation did not use out-of-sample errors to choose an order. The final fit below is a new fit of that already frozen full-sample order.

At h=2 across all primary-window targets, SARIMA had MAE {fmt_level(sarima_primary['MAE'])}, RMSE {fmt_level(sarima_primary['RMSE'])}, and MASE {fmt_mase(sarima_primary['MASE'])}. It beat each simple benchmark on all three metrics. The fixed ETS comparator also beat the simple benchmarks, but its h=2 all-target MAE, RMSE, and MASE were higher than SARIMA's.

*Table 1. Primary-window h=2 performance. MAE and RMSE are in 亿元; MASE is unit-free. Source: frozen Stage 3-6B metric tables.*

{historical_table}

## 8 Lunar New Year Incremental Value

Stage 5 adds one Q1-only, training-centered CNY timing regressor to the order selected by Stage 4 at each origin. On the primary h=2 all-target endpoint, the changes for SARIMA+CNY minus SARIMA were MAE {fmt_change(lower_cny_mae)}, RMSE {fmt_change(lower_cny_rmse)}, and MASE {fmt_change(lower_cny_mase, 6)}. The absolute reductions are small. In the {cny_q1['n']}-target Q1 subgroup, MAE changes by {fmt_change(fnum(cny_q1['MAE']) - fnum(sarima_q1['MAE']))}, RMSE by {fmt_change(fnum(cny_q1['RMSE']) - fnum(sarima_q1['RMSE']))}, and MASE by {fmt_change(fnum(cny_q1['MASE']) - fnum(sarima_q1['MASE']), 6)}; the subgroup metrics are mixed.

*Table 2. CNY minus SARIMA loss changes; negative values favor SARIMA+CNY.*

{cny_table}

The separate full-sample CNY coefficient and AICc are descriptive fit statistics, not evidence of out-of-sample forecasting value. The fixed-protocol result does not support a material or consistently positive incremental forecasting improvement. It describes forecast accuracy under the tested quarterly specification and does not identify a causal calendar effect.

## 9 PMI Extension

Stage 6A tests only lag-2 quarterly PMI, centered at 50, so both forecast-horizon inputs are available at the origin. Across {pmi_result['n']} h=2 all-target observations in the secondary window, the SARIMA+PMI model has higher MAE by {lower_pmi_mae:.2f}%, higher RMSE by {lower_pmi_rmse:.2f}%, and higher MASE by {lower_pmi_mase:.2f}% than SARIMA-common. The extension therefore does not improve this secondary endpoint.

*Table 3. Stage 6A h=2 all-target results on the separate secondary window.*

{pmi_table}

## 10 ETS and Shock Robustness

Stage 6B adds only the fixed additive-error, damped-additive-trend, additive-seasonality ETS specification with period 4. Its primary-window h=2 metrics are worse than SARIMA but better than the three simple benchmarks. The sensitivity analysis changes only the scored target rows for the preset quarters 2020Q1 and 2022Q2; it does not delete training observations or refit models.

*Table 4. SARIMA+CNY minus SARIMA changes before and after excluding the preset shock targets.*

{shock_table}

With both targets retained, CNY has a very small reduction on each aggregate metric. After excluding them, its MAE and MASE changes become small increases while the RMSE change remains a reduction. The interpretation remains mixed, and the shock-excluded aggregation does not replace the primary endpoint.

## 11 Final Forecast for 2026Q3 and 2026Q4

The final primary forecaster is the plain {order_label}, refit on {diag['training_start']}-{diag['training_end']} ({diag['n']} observations) and forecast from origin 2026Q2. It contains no CNY or PMI exogenous variables and uses no observations after 2026Q2. The conditional-mean point forecast is exp(mu + 0.5 v). Prediction intervals use exp(mu plus or minus the normal quantile times sqrt(v)); the 0.5 v adjustment is not applied to interval quantiles. The intervals represent forecast-distribution uncertainty conditional on fitted parameters and do not include parameter-estimation uncertainty.

![Final SARIMA forecast with prediction intervals](../outputs/figures/final_forecast.png)

*Figure 2. Observed retail through 2026Q2 and the h=1/h=2 primary forecast. Amounts are in RMB 100 million (亿元).*

*Table 5. Final primary forecasts from origin 2026Q2. Values and intervals are in 亿元.*

{forecast_markdown_table}

The supplementary comparison refits SARIMA+CNY with the Stage 5 CNY definition and uses the Stage 6B fixed ETS specification. The Q3 and Q4 CNY regressors equal zero because both targets are outside Q1. These forecasts are supplementary; they are not ensembled and do not determine the primary model.

*Table 6. Supplementary level forecasts; they do not alter the primary choice.*

{supplementary_markdown_table}

{diagnostic_note}

## 12 Limitations

The quarterly sample is modest and covers multiple growth regimes and unusual periods. Retail sales are nominal; these forecasts do not estimate real consumption growth. The lognormal intervals condition on fitted model parameters, and model-order uncertainty is not included. The CNY analysis tests one prespecified encoding and has limited Q1 targets; PMI uses one lag and a shorter secondary window without historical release-vintage adjustment. Shock sensitivity covers two prespecified target quarters only. None of these analyses establishes causality or guarantees forecast accuracy beyond the stated sample and model definitions.

## 13 Conclusion

Quarterly seasonality and a strong SARIMA baseline explain the main forecasting result under the frozen evaluation protocol. CNY timing does not show material and consistently positive out-of-sample improvement beyond that baseline, the fixed lag-2 PMI extension does not improve its separate h=2 comparison window, and fixed ETS remains behind SARIMA on the primary endpoint. The final forecasts therefore use the parsimonious Stage 4 full-sample SARIMA order, with the origin, training sample, point forecast, and prediction intervals reported explicitly above.

## 14 Reproducibility and Traceability

The forecast and report can be rebuilt with `python scripts/modeling/final_forecast.py` and `python scripts/reporting/build_final_report.py`. Run the repository test suite with `pytest -q` and compile the Python sources with `python -m compileall scripts tests`. Core numeric sources are `outputs/diagnostics/stage2_summary.json`, the Stage 3-6B tables under `outputs/tables/`, `outputs/forecasts/final_forecast.csv`, and `outputs/tables/final_model_diagnostics.csv`. The file `docs/stage7_protected_artifact_sha256.json` records the hashes used to verify that Stage 0-6B data and outputs remain unchanged.
"""
    return body


def _clean_inline(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", text)
    return text.replace("**", "").replace("`", "").replace("*", "")


def _set_run_font(run: Any, *, name: str = "Arial", size: float = 10.5, bold: bool | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold


def _shade_cell(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_margins(cell: Any, top: int = 95, start: int = 105, bottom: int = 95, end: int = 105) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_borders(table: Any) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "5")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D9D9D9")


def _repeat_table_header(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _no_row_split(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def _page_number(paragraph: Any) -> None:
    run = paragraph.add_run()
    _set_run_font(run, size=8.5)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, text, end):
        run._r.append(element)


def _add_word_table(document: Document, rows: list[list[str]], *, keep_together: bool = False) -> None:
    if not rows:
        return
    headers = rows[0]
    table = document.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    header = table.rows[0]
    _repeat_table_header(header)
    _no_row_split(header)

    width_map: dict[tuple[str, ...], list[float]] = {
        ("Model", "N", "MAE", "RMSE", "MASE"): [2.30, 0.55, 1.25, 1.25, 1.20],
        ("Comparison", "N", "Change in MAE", "Change in RMSE", "Change in MASE"): [1.95, 0.50, 1.45, 1.45, 1.35],
        ("Score sample", "N", "Change in MAE", "Change in RMSE", "Change in MASE"): [1.95, 0.50, 1.45, 1.45, 1.35],
        ("Target", "Horizon", "Median", "Conditional mean point forecast", "80% prediction interval", "95% prediction interval"): [0.70, 0.58, 0.95, 1.34, 1.56, 1.56],
        ("Model", "2026Q3 forecast", "2026Q4 forecast", "Forecast convention"): [1.05, 1.15, 1.15, 3.0],
    }
    widths = width_map.get(tuple(headers), [6.60 / len(headers)] * len(headers))
    for row_index, values in enumerate(rows):
        cells = header.cells if row_index == 0 else table.add_row().cells
        _no_row_split(table.rows[row_index])
        for column, raw_text in enumerate(values):
            cell = cells[column]
            cell.text = ""
            cell.width = Inches(widths[column])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)
            _shade_cell(cell, "DCE8F1" if row_index == 0 else ("F5F8FA" if row_index % 2 == 0 else "FFFFFF"))
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(1.5)
            paragraph.paragraph_format.space_before = Pt(1.5)
            if row_index == 0:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif column > 0 and (re.fullmatch(r"[0-9.,+\-\[\] ]+", raw_text) or headers[column] in ("N", "MAE", "RMSE", "MASE")):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(raw_text)
            _set_run_font(run, size=8.0 if len(headers) >= 6 else 8.5, bold=(row_index == 0))
    if keep_together:
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_together = True
                    paragraph.paragraph_format.keep_with_next = row_index < len(table.rows) - 1
    _set_table_borders(table)
    document.add_paragraph().paragraph_format.space_after = Pt(3)


def _add_markdown_to_docx(markdown: str) -> Document:
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for style_name, size in (("Heading 1", 15), ("Heading 2", 12)):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("STA4003 Time Series Project  |  Page ")
    _page_number(footer)
    for run in footer.runs:
        _set_run_font(run, size=8.5)
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    document.settings._element.append(update_fields)

    lines = markdown.splitlines()
    index = 0
    title_page = True
    first_title = True
    keep_table_three_together = False
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line == "<!-- PAGE_BREAK -->":
            document.add_page_break()
            title_page = False
            index += 1
            continue
        if line.startswith("# "):
            paragraph = document.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(120 if first_title else 6)
            paragraph.paragraph_format.space_after = Pt(16)
            run = paragraph.add_run(_clean_inline(line[2:]))
            _set_run_font(run, size=25, bold=True)
            first_title = False
            index += 1
            continue
        if line.startswith("## "):
            paragraph = document.add_paragraph(style="Heading 1")
            run = paragraph.add_run(_clean_inline(line[3:]))
            _set_run_font(run, size=15, bold=True)
            index += 1
            continue
        if line.startswith("### "):
            paragraph = document.add_paragraph(style="Heading 2")
            run = paragraph.add_run(_clean_inline(line[4:]))
            _set_run_font(run, size=12, bold=True)
            index += 1
            continue
        if line.startswith("!["):
            image_match = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line)
            if image_match:
                alt, relative = image_match.groups()
                path = (MARKDOWN_PATH.parent / relative).resolve()
                if not path.is_file():
                    raise FileNotFoundError(f"Report figure is missing: {path}")
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.keep_together = True
                paragraph.paragraph_format.keep_with_next = True
                shape = paragraph.add_run().add_picture(str(path), width=Inches(6.55))
                shape._inline.docPr.set("descr", alt)
                shape._inline.docPr.set("title", alt)
                index += 1
                continue
        if line.startswith("|"):
            table_rows: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                    table_rows.append([_clean_inline(cell) for cell in cells])
                index += 1
            _add_word_table(document, table_rows, keep_together=keep_table_three_together)
            keep_table_three_together = False
            continue
        if line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(3)
            run = paragraph.add_run(_clean_inline(line[2:]))
            _set_run_font(run, size=10)
            index += 1
            continue
        if line.startswith("*") and line.endswith("*") and len(line) > 2:
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.space_after = Pt(7)
            run = paragraph.add_run(_clean_inline(line))
            _set_run_font(run, size=8.8)
            run.italic = True
            if _clean_inline(line).startswith("Table 3."):
                paragraph.paragraph_format.keep_with_next = True
                keep_table_three_together = True
            index += 1
            continue
        paragraph = document.add_paragraph()
        if title_page:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(8)
            paragraph.paragraph_format.keep_together = True
            run = paragraph.add_run(_clean_inline(line).replace("  ", " "))
            _set_run_font(run, size=11, bold=(line.startswith("**Course**") or line.startswith("**Report**")))
        else:
            paragraph.paragraph_format.keep_together = True
            run = paragraph.add_run(_clean_inline(line))
            _set_run_font(run, size=10.5)
        index += 1

    document.core_properties.title = "Forecasting Quarterly Retail Sales in China"
    document.core_properties.subject = "STA4003 Time Series Project final report"
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.created = datetime(2026, 10, 3, 0, 0, 0)
    document.core_properties.modified = datetime(2026, 10, 3, 0, 0, 0)
    return document


def _word_converter_path() -> Path | None:
    return next((path for path in WORD_PATHS if path.is_file()), None)


def _canonicalize_docx_timestamps(docx_path: Path) -> None:
    """Fix ZIP member timestamps so identical report content has identical bytes."""
    timestamp = (2020, 1, 1, 0, 0, 0)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{docx_path.stem}-", suffix=".tmp.docx", dir=docx_path.parent
    )
    os.close(descriptor)
    temporary_path = Path(temporary_name)
    try:
        with zipfile.ZipFile(docx_path, "r") as source:
            with zipfile.ZipFile(temporary_path, "w") as destination:
                for original in source.infolist():
                    normalized = zipfile.ZipInfo(original.filename, date_time=timestamp)
                    normalized.compress_type = original.compress_type
                    normalized.comment = original.comment
                    normalized.create_system = original.create_system
                    normalized.external_attr = original.external_attr
                    normalized.extra = original.extra
                    normalized.internal_attr = original.internal_attr
                    destination.writestr(normalized, source.read(original.filename))
        os.replace(temporary_path, docx_path)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()


def _export_pdf_with_word(docx_path: Path, pdf_path: Path) -> None:
    """Use the already installed, hidden local Word COM server for PDF export."""
    if _word_converter_path() is None:
        return
    powershell = shutil.which("powershell.exe") or r"C:\Windows\System32\WindowsPowerShell\v1.0\powershell.exe"
    if not Path(powershell).is_file():
        raise RuntimeError("Microsoft Word is installed, but Windows PowerShell is unavailable for local PDF export")

    ps_source = r"""
param([string]$DocxPath, [string]$PdfPath)
$ErrorActionPreference = 'Stop'
$word = $null
$document = $null
try {
    $word = New-Object -ComObject Word.Application
    $word.Visible = $false
    $word.DisplayAlerts = 0
    $document = $word.Documents.Open($DocxPath, $false, $true, $false)
    $document.Repaginate()
    foreach ($section in $document.Sections) {
        foreach ($footer in $section.Footers) { $footer.Range.Fields.Update() | Out-Null }
    }
    $document.ExportAsFixedFormat($PdfPath, 17)
    $document.Close(0)
    [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($document)
    $document = $null
}
finally {
    if ($document -ne $null) {
        $document.Close(0)
        [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($document)
    }
    if ($word -ne $null) {
        $word.Quit(0)
        [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($word)
    }
}
if (-not (Test-Path -LiteralPath $PdfPath -PathType Leaf)) { throw 'Word did not create the requested PDF.' }
if ((Get-Item -LiteralPath $PdfPath).Length -lt 1000) { throw 'Word exported an empty or incomplete PDF.' }
"Exported $PdfPath"
"""
    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_pdf = pdf_path.with_name(f".{pdf_path.stem}.tmp.pdf")
    if temporary_pdf.exists():
        temporary_pdf.unlink()
    try:
        with tempfile.TemporaryDirectory(prefix="sta4003-word-export-") as temp_dir:
            ps_path = Path(temp_dir) / "export_docx_to_pdf.ps1"
            ps_path.write_text(ps_source, encoding="utf-8-sig")
            try:
                subprocess.run(
                    [powershell, "-NoProfile", "-NonInteractive", "-File", str(ps_path), str(docx_path), str(temporary_pdf)],
                    check=True,
                    capture_output=True,
                    text=True,
                    encoding="utf-8",
                    errors="replace",
                )
            except subprocess.CalledProcessError as error:
                details = " ".join(part.strip() for part in (error.stdout or "", error.stderr or "") if part.strip())
                raise RuntimeError(f"Word PDF export failed with exit {error.returncode}: {details}") from error
        if temporary_pdf.read_bytes()[:5] != b"%PDF-":
            raise RuntimeError("Word's PDF output does not have a valid PDF header")
        os.replace(temporary_pdf, pdf_path)
    finally:
        if temporary_pdf.exists():
            temporary_pdf.unlink()


def _sync_pdf_status(conversion_error: str | None = None) -> None:
    summary = json.loads(STAGE7_SUMMARY_PATH.read_text(encoding="utf-8"))
    converter = _word_converter_path()
    generated = converter is not None and PDF_PATH.is_file() and PDF_PATH.stat().st_size > 0
    if generated and conversion_error is None:
        reason = None
    elif conversion_error and re.search(r"0x80070520", conversion_error, re.IGNORECASE):
        reason = (
            "Microsoft Word is installed, but COM PDF export could not start "
            "because this process has no active Windows logon session "
            "(HRESULT 0x80070520)."
        )
    elif conversion_error:
        reason = "A local DOCX-to-PDF conversion attempt failed; no PDF was generated."
    else:
        reason = "No reliable local DOCX-to-PDF converter was available."
    summary["pdf_status"] = {
        "converter_installed": converter is not None,
        "converter_available": generated and conversion_error is None,
        "runtime_tested": True,
        "converter": "Microsoft Word" if converter is not None else None,
        "generated": generated,
        "output": "outputs/reports/STA4003_Final_Report.pdf" if generated else None,
        "reason": reason,
    }
    artifacts = [
        item for item in summary["generated_artifacts"]
        if item != "outputs/reports/STA4003_Final_Report.pdf"
    ]
    if generated:
        artifacts.append("outputs/reports/STA4003_Final_Report.pdf")
    summary["generated_artifacts"] = artifacts
    STAGE7_SUMMARY_PATH.write_text(
        json.dumps(summary, indent=2, ensure_ascii=False, allow_nan=False, sort_keys=True) + "\n",
        encoding="utf-8",
    )


def build_final_report() -> tuple[Path, Path]:
    markdown = build_markdown()
    MARKDOWN_PATH.parent.mkdir(parents=True, exist_ok=True)
    MARKDOWN_PATH.write_text(markdown, encoding="utf-8", newline="\n")
    DOCX_PATH.parent.mkdir(parents=True, exist_ok=True)
    document = _add_markdown_to_docx(markdown)
    document.save(DOCX_PATH)
    _canonicalize_docx_timestamps(DOCX_PATH)
    conversion_error: str | None = None
    try:
        _export_pdf_with_word(DOCX_PATH, PDF_PATH)
    except (OSError, RuntimeError, subprocess.SubprocessError) as error:
        conversion_error = str(error)
    _sync_pdf_status(conversion_error)
    return MARKDOWN_PATH, DOCX_PATH


def main() -> None:
    markdown_path, docx_path = build_final_report()
    print(f"Wrote {markdown_path.relative_to(PROJECT_ROOT)}")
    print(f"Wrote {docx_path.relative_to(PROJECT_ROOT)}")
    if PDF_PATH.is_file():
        print(f"Wrote {PDF_PATH.relative_to(PROJECT_ROOT)} using local Microsoft Word")
    else:
        summary = json.loads(STAGE7_SUMMARY_PATH.read_text(encoding="utf-8"))
        print(f"PDF not generated: {summary['pdf_status']['reason']}")


if __name__ == "__main__":
    main()
